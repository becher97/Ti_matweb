from flask import Flask, request, jsonify, render_template
from io import BytesIO
from datetime import datetime
import shutil
import sqlite3
import csv
import os
# Optional: used for reading .xls
try:
    import pandas as pd  # type: ignore
except Exception:
    pd = None

app = Flask(__name__)
DB_FILE = "materials.db"
CSV_FILE = "materials.csv"
XLS_FILE = "results-csv.xls"  # legacy .xls
XLSX_FILE = "results-csv.xlsx"  # preferred .xlsx
MAIN_DB = DB_FILE
USER_DB = "user.db"
# Accept both variants for name column
NAME_CANDIDATES = ["合金成分", "合金成份", "name", "Name"]

# Active database path (can be switched at runtime)
CURRENT_DB = MAIN_DB
# Prefer .xlsx if present
try:
    if os.path.exists(XLSX_FILE):
        XLS_FILE = XLSX_FILE
except Exception:
    pass

# 每次启动时根据 CSV 更新数据库
def update_db_from_csv():
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)  # 删除旧数据库

    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()

    # 读取 CSV
    with open(CSV_FILE, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames

        # 创建表
        columns_sql = []
        for field in fieldnames:
            field_escaped = f'"{field}"'  # 用双引号包裹列名，支持单位或特殊字符
            if field.lower() == "id":
                columns_sql.append(f"{field_escaped} INTEGER PRIMARY KEY")
            elif field.lower() in ["name", "category"]:
                columns_sql.append(f"{field_escaped} TEXT")
            else:
                columns_sql.append(f"{field_escaped} REAL")
        c.execute(f"CREATE TABLE materials ({', '.join(columns_sql)})")

        # 插入数据
        placeholders = ", ".join("?" for _ in fieldnames)
        fields_escaped = [f'"{f}"' for f in fieldnames]
        for row in reader:
            values = [row[f] if row[f] != "" else None for f in fieldnames]
            c.execute(f'INSERT INTO materials ({", ".join(fields_escaped)}) VALUES ({placeholders})', values)

    conn.commit()
    conn.close()

# 获取可筛选属性（数值列）
@app.route("/properties", methods=["GET"])
def properties():
    conn = sqlite3.connect(CURRENT_DB)
    c = conn.cursor()
    c.execute("PRAGMA table_info(materials)")
    fields = c.fetchall()
    all_fields = [f[1] for f in fields if f[1].lower() not in ["id", "name", "category"]]

    result = {}
    for field in all_fields:
        c.execute(f'SELECT "{field}" FROM materials')
        values = [v[0] for v in c.fetchall()]
        # 检查是否全为数值或空值
        numeric_values = []
        is_numeric = True
        for v in values:
            if v is None or v == "":
                continue
            try:
                numeric_values.append(float(v))
            except:
                is_numeric = False
                break
        if is_numeric and numeric_values:
            result[field] = {"min": min(numeric_values), "max": max(numeric_values)}
    conn.close()
    return jsonify(result)

# 初始化数据库：优先从 XLS 构建；若无则回退至 CSV
def init_db():
    if os.path.exists(DB_FILE):
        os.remove(DB_FILE)

    conn = sqlite3.connect(MAIN_DB)
    c = conn.cursor()
    try:
        if os.path.exists(XLS_FILE):
            if pd is None:
                raise RuntimeError("需要安装 pandas 与 xlrd 才能读取 .xls 文件")
            df = pd.read_excel(XLS_FILE)
            columns = [str(col) for col in df.columns]
            cols_sql = ", ".join([f'"{col}" TEXT' for col in columns])
            c.execute(f"CREATE TABLE materials ({cols_sql})")
            placeholders = ", ".join(["?" for _ in columns])
            fields_escaped = ", ".join([f'"{col}"' for col in columns])
            for _, row in df.iterrows():
                vals = []
                for v in row.tolist():
                    if pd is not None and hasattr(pd, 'isna') and pd.isna(v):
                        vals.append(None)
                    else:
                        vals.append(v)
                c.execute(f"INSERT INTO materials ({fields_escaped}) VALUES ({placeholders})", vals)
        else:
            with open(CSV_FILE, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                fieldnames = reader.fieldnames or []
                if not fieldnames:
                    raise RuntimeError("CSV 表头为空")
                columns_sql = []
                for field in fieldnames:
                    field_escaped = f'"{field}"'
                    if field.lower() == "id":
                        columns_sql.append(f"{field_escaped} INTEGER PRIMARY KEY")
                    elif field.lower() in ["name", "category"]:
                        columns_sql.append(f"{field_escaped} TEXT")
                    else:
                        columns_sql.append(f"{field_escaped} REAL")
                c.execute(f"CREATE TABLE materials ({', '.join(columns_sql)})")
                placeholders = ", ".join("?" for _ in fieldnames)
                fields_escaped = [f'"{f}"' for f in fieldnames]
                for row in reader:
                    values = [row[f] if row[f] != "" else None for f in fieldnames]
                    c.execute(f'INSERT INTO materials ({", ".join(fields_escaped)}) VALUES ({placeholders})', values)
    finally:
        conn.commit()
        conn.close()

def _choose_name_column_from(conn):
    c = conn.cursor()
    c.execute("PRAGMA table_info(materials)")
    cols = [r[1] for r in c.fetchall()]
    for nm in NAME_CANDIDATES:
        if nm in cols:
            return nm
    return None

# 搜索接口
@app.route("/search", methods=["POST"])
def search():
    conditions = request.json.get("conditions", [])
    conn = sqlite3.connect(CURRENT_DB)
    c = conn.cursor()
    # 获取实际列
    c.execute("PRAGMA table_info(materials)")
    table_cols = [f[1] for f in c.fetchall()]
    # 选择列：rowid 作为 id，name 使用合金成分（若存在，否则空串）
    select_list = ["rowid AS id"]
    name_col = None
    for cand in NAME_CANDIDATES:
        if cand in table_cols:
            name_col = cand
            break
    if name_col:
        select_list.append(f'"{name_col}" AS name')
    else:
        select_list.append("'' AS name")
    # 追加实际列
    select_list.extend([f'"{cname}"' for cname in table_cols])

    query = f"SELECT {', '.join(select_list)} FROM materials WHERE 1=1"
    params = []
    for cond in conditions:
        prop = cond.get("property")
        if not prop:
            continue
        min_val = cond.get("min")
        max_val = cond.get("max")
        query += f' AND CAST("{prop}" AS REAL) BETWEEN ? AND ?'
        params.extend([min_val, max_val])

    c.execute(query, params)
    rows = c.fetchall()
    columns = ["id", "name", *table_cols]
    conn.close()

    results = [dict(zip(columns, r)) for r in rows]
    return jsonify({"columns": columns, "data": results})

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/stats", methods=["GET"])
def stats():
    try:
        conn = sqlite3.connect(CURRENT_DB)
        c = conn.cursor()
        c.execute("SELECT COUNT(*) FROM materials")
        total = c.fetchone()[0]
        conn.close()
    except Exception:
        total = None
    return jsonify({"total": total})

@app.route("/shutdown", methods=["POST"])
def shutdown():
    import os
    func = request.environ.get("werkzeug.server.shutdown")
    if func:
        func()
        return "Server shutting down..."
    else:
        os._exit(0)

# 获取单条记录（通过 rowid）
@app.route("/item/<int:item_id>", methods=["GET"])
def get_item(item_id: int):
    conn = sqlite3.connect(CURRENT_DB)
    c = conn.cursor()
    # 获取列
    c.execute("PRAGMA table_info(materials)")
    table_cols = [f[1] for f in c.fetchall()]
    # 构建选择列：rowid 作为 id，"合金成分" 作为 name（若存在）
    select_list = ["rowid AS id"]
    name_col = None
    for cand in NAME_CANDIDATES:
        if cand in table_cols:
            name_col = cand
            break
    if name_col:
        select_list.append(f'"{name_col}" AS name')
    else:
        select_list.append("'' AS name")
    select_list.extend([f'"{cname}"' for cname in table_cols])

    query = f"SELECT {', '.join(select_list)} FROM materials WHERE rowid = ?"
    c.execute(query, (item_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "not found"}), 404
    columns = ["id", "name", *table_cols]
    return jsonify({"data": dict(zip(columns, row))})

# Database switching and upload
@app.route("/db/options", methods=["GET"])
def db_options():
    opts = []
    opts.append({"key": "main", "label": "主数据库", "path": MAIN_DB, "exists": os.path.exists(MAIN_DB)})
    opts.append({"key": "user", "label": "用户数据库", "path": USER_DB, "exists": os.path.exists(USER_DB)})
    cur = "main" if CURRENT_DB == MAIN_DB else ("user" if CURRENT_DB == USER_DB else "temp")
    return jsonify({"current": cur, "options": opts})

@app.route("/db/select", methods=["POST"])
def db_select():
    global CURRENT_DB
    key = (request.json or {}).get("db")
    if key == "main":
        CURRENT_DB = MAIN_DB
        return jsonify({"ok": True, "current": "main"})
    elif key == "user":
        # If user db not exists, create by copying main
        if not os.path.exists(USER_DB):
            if not os.path.exists(MAIN_DB):
                return jsonify({"error": "main db missing"}), 500
            shutil.copy(MAIN_DB, USER_DB)
        CURRENT_DB = USER_DB
        return jsonify({"ok": True, "current": "user"})
    return jsonify({"error": "unknown db"}), 400

@app.route("/db/upload", methods=["POST"])
def db_upload():
    global CURRENT_DB
    if 'file' not in request.files:
        return jsonify({"error": "缺少文件"}), 400
    f = request.files['file']
    if not f.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({"error": "仅支持 .xlsx 或 .xls"}), 400
    # Save to temp
    os.makedirs('uploads', exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d-%H%M%S')
    tmp_xlsx = os.path.join('uploads', f'upload-{ts}.xlsx')
    f.save(tmp_xlsx)
    # Build sqlite db from xlsx
    tmp_db = os.path.join('uploads', f'upload-{ts}.db')
    try:
        if pd is None:
            return jsonify({"error": "需要安装 pandas 与 openpyxl/xlrd"}), 500
        try:
            df = pd.read_excel(tmp_xlsx, engine='openpyxl')
        except Exception:
            # fallback try xlrd
            df = pd.read_excel(tmp_xlsx, engine='xlrd')
        conn = sqlite3.connect(tmp_db)
        c = conn.cursor()
        columns = [str(col) for col in df.columns]
        cols_sql = ", ".join([f'"{col}" TEXT' for col in columns])
        c.execute(f"CREATE TABLE materials ({cols_sql})")
        placeholders = ", ".join(["?" for _ in columns])
        fields_escaped = ", ".join([f'"{col}"' for col in columns])
        for _, row in df.iterrows():
            vals = []
            for v in row.tolist():
                try:
                    is_na = pd.isna(v)
                except Exception:
                    is_na = False
                vals.append(None if is_na else v)
            c.execute(f"INSERT INTO materials ({fields_escaped}) VALUES ({placeholders})", vals)
        conn.commit()
        conn.close()
    except Exception as e:
        try:
            os.remove(tmp_db)
        except Exception:
            pass
        return jsonify({"error": f"构建数据库失败: {e}"}), 500
    try:
        if os.path.exists(USER_DB):
            try:
                os.remove(USER_DB)
            except Exception:
                pass
        shutil.copyfile(tmp_db, USER_DB)
        CURRENT_DB = USER_DB
        return jsonify({"ok": True, "current": "user", "db_path": USER_DB})
    except Exception as e:
        CURRENT_DB = tmp_db
        return jsonify({"ok": True, "current": "temp", "db_path": tmp_db, "warn": f"user db replace failed: {e}"})

@app.route("/export", methods=["GET"])
def export_candidates():
    ids_param = request.args.get("ids", "").strip()
    if not ids_param:
        return jsonify({"error": "missing ids"}), 400
    try:
        ids = [int(x) for x in ids_param.split(",") if x.strip() != ""]
    except Exception:
        return jsonify({"error": "invalid ids"}), 400
    if not ids:
        return jsonify({"error": "no ids"}), 400

    conn = sqlite3.connect(CURRENT_DB)
    c = conn.cursor()
    c.execute("PRAGMA table_info(materials)")
    table_cols = [f[1] for f in c.fetchall()]
    col_list = ", ".join([f'"{col}"' for col in table_cols])
    placeholders = ",".join(["?" for _ in ids])
    query = f"SELECT {col_list} FROM materials WHERE rowid IN ({placeholders}) ORDER BY rowid"
    c.execute(query, ids)
    rows = c.fetchall()
    conn.close()

    if pd is None:
        return jsonify({"error": "需要安装 pandas 与 openpyxl 才能导出 xlsx"}), 500
    try:
        import openpyxl  # noqa: F401
    except Exception:
        return jsonify({"error": "需要安装 openpyxl 才能导出 xlsx"}), 500

    try:
        df = pd.DataFrame(rows, columns=table_cols)
        bio = BytesIO()
        df.to_excel(bio, index=False, engine='openpyxl')
        bio.seek(0)
    except Exception as e:
        return jsonify({"error": f"导出失败: {e}"}), 500

    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"candidates-{ts}.xlsx"
    resp = app.response_class(bio.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp

if __name__ == "__main__":
    # 启动时更新数据库（优先从 XLS 读取）
    init_db()
    app.run(debug=True, use_reloader=False)
